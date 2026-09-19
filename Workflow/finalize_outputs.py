"""Format KiCad outputs inside a Prism-selected Jobset destination.

This does not run exports, create releases, invoke Git, or publish anything.
KiCad supplies the temporary destination through JOBSET_OUTPUT_WORK_PATH.
"""

from __future__ import annotations

import argparse
import csv
from datetime import datetime
import json
import os
from pathlib import Path
import re
import shutil
import sys
from tempfile import TemporaryDirectory


HEADERS = [
    "Line #", "Quantity", "Designator", "Description", "Manufacturer",
    "Manufacturer Part Number",
]


def filename_value(value: str) -> str:
    if not value or re.search(r'[<>:"/\\|?*\x00-\x1f]', value) or value.endswith((" ", ".")):
        raise ValueError(f"Not a portable filename value: {value!r}")
    return value


def create_bom(source: Path, target: Path, variables: dict, variant: str, published: datetime) -> None:
    from copy import copy
    import openpyxl
    from openpyxl.styles import Alignment, Font

    with source.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if reader.fieldnames != HEADERS:
            raise ValueError(f"Unexpected native BOM columns: {reader.fieldnames}")
        rows = list(reader)

    workbook = openpyxl.load_workbook(Path(__file__).with_name("BOM-template.xlsx"))
    sheet = workbook.active
    for column, factor in (("E", 1.20), ("F", 1.75)):
        dimension = sheet.column_dimensions[column]
        baseline = dimension.width
        if baseline is None:
            baseline = sheet.sheet_format.defaultColWidth
        dimension.width = (baseline if baseline is not None else 13.0) * factor
    if [sheet.cell(8, c).value for c in range(1, 7)] != HEADERS:
        raise ValueError("BOM template header changed; update field mapping before release")

    sheet["C2"] = variables["ProjectTitle"]
    sheet["C3"] = variables.get("ProjectPartNumber") or "N/A"
    sheet["C4"] = str(variables["ProjectPCBARevision"])
    sheet["C5"] = published.replace(tzinfo=None)
    sheet["C5"].number_format = "yyyy-mm-dd hh-mm"
    sheet["C6"] = variant
    # Revisions, identifiers and descriptions are text, never spreadsheet formulas.
    for address in ("C2", "C3", "C4", "C6"):
        sheet[address].data_type = "s"
        sheet[address].number_format = "@"

    for row_number, row in enumerate(rows, 9):
        for column, header in enumerate(HEADERS, 1):
            cell = sheet.cell(row_number, column)
            value = row[header]
            cell.value = int(value) if column <= 2 else value
            if column > 2:
                cell.data_type = "s"
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(vertical="top", wrap_text=True,
                                       horizontal="center" if column <= 2 else "left")
            cell.border = copy(sheet.cell(8, column).border)
        # Leave wrapped item rows eligible for the spreadsheet reader's AutoFit.
        sheet.row_dimensions[row_number].height = None

    for col in range(1, 7):
        alignment = copy(sheet.cell(8, col).alignment)
        alignment.wrap_text = True
        sheet.cell(8, col).alignment = alignment
    sheet.row_dimensions[8].height = max(sheet.row_dimensions[8].height or 0, 30)
    sheet.print_area = f"A1:F{max(8, 8 + len(rows))}"
    sheet.print_title_rows = "1:8"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    workbook.save(target)


def merge_assembly(top: Path, bottom: Path, target: Path, title: str) -> None:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for source, label in ((top, "Top"), (bottom, "Bottom")):
        reader = PdfReader(source)
        if len(reader.pages) != 1:
            raise ValueError(f"Expected one {label} drawing page, got {len(reader.pages)}")
        writer.append(reader, outline_item=label, import_outline=False)
    writer.add_metadata({"/Title": title, "/Subject": "PCBA assembly: top and mirrored bottom"})
    with target.open("wb") as stream:
        writer.write(stream)


def jobset_work_root() -> Path:
    project_root = Path(__file__).resolve().parent.parent
    temporary = os.environ.get("JOBSET_OUTPUT_WORK_PATH")
    if not temporary:
        raise RuntimeError("Run this formatter only through Outputs.kicad_jobset in KiCad/Prism")
    root = Path(temporary).resolve(strict=True)
    if root == project_root or root in project_root.parents or project_root in root.parents:
        raise RuntimeError("Formatter output must be KiCad's temporary destination, outside source")
    return root


def prepare_assembly_worksheets() -> None:
    root = jobset_work_root()
    project_root = Path(__file__).resolve().parent.parent
    library = Path(os.environ.get("KICAD_LIB_ROOT") or project_root.parent / "kicad-lib")
    source = library / "drawing-sheets" / "alex-generic-pcba.kicad_wks"
    if not source.is_file():
        raise FileNotFoundError(f"Canonical PCBA worksheet missing: {source}. "
                                "Set KICAD_LIB_ROOT to the shared kicad-lib checkout.")
    # Read once so both views use the same snapshot, even if the library is edited.
    # Bytes preserve encoding, BOM and line endings; all other variables stay native.
    content = source.read_bytes()
    if b"${#}" not in content or b"${##}" not in content:
        raise ValueError("Canonical worksheet must contain ${#} and ${##} page tokens")
    work = root / "_work"
    work.mkdir(exist_ok=True)
    for page, view in ((1, "top"), (2, "bottom")):
        target = work / f"assembly-{view}.kicad_wks"
        # Exclusive creation prevents following a pre-existing target/symlink.
        with target.open("xb") as stream:
            stream.write(content.replace(b"${##}", b"2").replace(b"${#}", str(page).encode()))
    print(f"Prepared assembly worksheets from {source.resolve()}")


def finalize(kind: str, variant_name: str = "") -> list[Path]:
    project_root = Path(__file__).resolve().parent.parent
    root = jobset_work_root()
    projects = list(project_root.glob("*.kicad_pro"))
    if len(projects) != 1:
        raise RuntimeError("Expected exactly one KiCad project alongside the canonical Jobset")
    variables = json.loads(projects[0].read_text(encoding="utf-8"))["text_variables"]
    title = filename_value(variables["ProjectTitle"])
    revision = filename_value(str(variables["ProjectPCBRevision" if kind == "fabrication" else "ProjectPCBARevision"]))
    variant = filename_value(variant_name or "No Variant")
    published = datetime.now().astimezone().replace(second=0, microsecond=0)
    suffix = f" - {variant} ({published:%Y-%m-%d %H-%M})"
    prefix = f"{title} R{revision} "
    work = root / "_work"

    if kind == "fabrication":
        mapping = {"fabrication.pdf": ("FAB", ".pdf"), "envelope.step": ("STEP", ".step")}
        inputs = list(mapping)
        destination = root
    else:
        mapping = {"schematic.pdf": ("SCH", ".pdf"), "assembly-3d.pdf": ("3DPCB", ".pdf"),
                   "positions-all-pos.csv": ("Pick Place", ".csv")}
        inputs = list(mapping) + ["assembly-top.pdf", "assembly-bottom.pdf", "bom.csv"]
        destination = root / variant

    for name in inputs:
        source = work / name
        if source.is_symlink() or not source.is_file() or not source.stat().st_size:
            raise RuntimeError(f"Missing or invalid native output: {name}")

    # Build every formatted file before placing any of them in the destination.
    with TemporaryDirectory(prefix="format-", dir=root) as staging_path:
        staging = Path(staging_path)
        for name, (file_type, extension) in mapping.items():
            shutil.copyfile(work / name, staging / (prefix + file_type + suffix + extension))
        if kind == "assembly":
            merge_assembly(work / "assembly-top.pdf", work / "assembly-bottom.pdf",
                           staging / (prefix + "ASY" + suffix + ".pdf"), prefix + "ASY" + suffix)
            create_bom(work / "bom.csv", staging / (prefix + "BOM" + suffix + ".xlsx"),
                       variables, variant, published)
        destination.mkdir(parents=True, exist_ok=True)
        result = []
        for source in sorted(staging.iterdir()):
            target = destination / source.name
            source.replace(target)
            result.append(target)

    for name in inputs:
        (work / name).unlink()
    if kind == "assembly":
        for view in ("top", "bottom"):
            (work / f"assembly-{view}.kicad_wks").unlink(missing_ok=True)
    work.rmdir()
    for path in result:
        print(f"Created {path.relative_to(root).as_posix()}")
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--kind", choices=("fabrication", "assembly", "prepare-assembly"), required=True)
    parser.add_argument("--variant", default="", help="Empty selects the base design, labelled No Variant")
    args = parser.parse_args()
    try:
        if args.kind == "prepare-assembly":
            prepare_assembly_worksheets()
        else:
            finalize(args.kind, args.variant)
    except Exception as error:
        print(f"Output formatting failed: {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
