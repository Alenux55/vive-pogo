"""Tests for the formatting boundary; exports are validated by running the Jobset."""

import csv
from datetime import datetime
import json
import os
from pathlib import Path
import struct
import tempfile
import unittest
from unittest.mock import patch

import openpyxl
from pypdf import PdfReader, PdfWriter

import finalize_outputs as formatter


class FormattingTests(unittest.TestCase):
    @staticmethod
    def _u3d_material(name, diffuse):
        encoded = name.encode()
        payload = (struct.pack("<H", len(encoded)) + encoded + struct.pack("<I", 0x36)
                   + struct.pack("<3f", 0, 0, 0) + struct.pack("<3f", *diffuse)
                   + struct.pack("<3f", 0.2, 0.2, 0.2) + struct.pack("<3f", 0, 0, 0)
                   + struct.pack("<2f", 0.32, 1.0))
        return (struct.pack("<III", 0xFFFFFF54, len(payload), 0) + payload
                + bytes((-len(payload)) % 4))

    def test_recolors_only_named_u3d_pad_material(self):
        copper = (0.7, 0.61, 0.0)
        original = (self._u3d_material("m_Copper_0", copper)
                    + self._u3d_material("m_Pads_0", (0.5, 0.5, 0.5))
                    + self._u3d_material("m_Component_0", (0.5, 0.5, 0.5)))
        changed = formatter.recolor_u3d_pads(original)
        self.assertNotEqual(changed, original)
        self.assertEqual(changed.count(struct.pack("<3f", *copper)), 2)
        self.assertEqual(changed.count(struct.pack("<3f", 0.5, 0.5, 0.5)), 1)
        self.assertEqual(formatter.recolor_u3d_pads(changed), changed)

    def test_recolor_rejects_unexpected_u3d_materials(self):
        data = self._u3d_material("m_Copper_0", (0.7, 0.61, 0.0))
        with self.assertRaisesRegex(ValueError, "m_Pads_0"):
            formatter.recolor_u3d_pads(data)

    def test_requires_jobset_context(self):
        with patch.dict(os.environ, {}, clear=True):
            with self.assertRaisesRegex(RuntimeError, "only through"):
                formatter.finalize("assembly")

    def test_assembly_page_count_follows_enabled_bottom_job(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "Outputs.kicad_jobset"
            base = {
                "jobs": [
                    {"id": "prepare", "settings": {"command": "tool --kind prepare-assembly"}},
                    {"id": "bottom", "settings": {"output_filename": "_work/assembly-bottom.pdf"}},
                ],
                "outputs": [{"only": ["prepare", "bottom"]}],
            }
            path.write_text(json.dumps(base))
            self.assertEqual(formatter.assembly_page_count(path), 2)
            base["outputs"][0]["only"].remove("bottom")
            path.write_text(json.dumps(base))
            self.assertEqual(formatter.assembly_page_count(path), 1)

    def test_rejects_path_characters(self):
        for value in ("../part", "part/name", "part:name", "", "part."):
            with self.subTest(value=value), self.assertRaises(ValueError):
                formatter.filename_value(value)

    def test_missing_export_does_not_publish_partial_package(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "_work").mkdir()
            (root / "_work/fabrication.pdf").write_bytes(b"native output")
            with patch.dict(os.environ, JOBSET_OUTPUT_WORK_PATH=directory):
                with self.assertRaisesRegex(RuntimeError, "envelope.step"):
                    formatter.finalize("fabrication")
            self.assertEqual(list(root.iterdir()), [root / "_work"])
            self.assertTrue((root / "_work/fabrication.pdf").exists())

    def test_template_metadata_types_and_literal_text(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "bom.csv"
            target = Path(directory) / "bom.xlsx"
            with source.open("w", newline="", encoding="utf-8") as stream:
                writer = csv.writer(stream)
                writer.writerow(formatter.HEADERS)
                writer.writerow([1, 2, "R1,R2", "=literal description", "Maker", "00123"])
            published = datetime(2026, 9, 18, 17, 5)
            formatter.create_bom(source, target,
                                 {"ProjectTitle": "Board", "ProjectPCBARevision": "1.0"},
                                 "No Variant", published)
            sheet = openpyxl.load_workbook(target).active
            self.assertEqual(sheet["C4"].value, "1.0")
            self.assertEqual(sheet["C4"].number_format, "@")
            self.assertEqual(sheet["C5"].value, published)
            self.assertEqual(sheet["C6"].value, "No Variant")
            self.assertEqual(sheet["A9"].value, 1)
            self.assertEqual(sheet["B9"].value, 2)
            self.assertEqual(sheet["D9"].data_type, "s")
            self.assertEqual(sheet["F9"].value, "00123")
            self.assertIn("A1:C1", str(sheet.merged_cells))
            template = openpyxl.load_workbook(Path(formatter.__file__).with_name("BOM-template.xlsx")).active
            for column, factor in (("E", 1.20), ("F", 1.75)):
                self.assertAlmostEqual(sheet.column_dimensions[column].width,
                                       template.column_dimensions[column].width * factor)
            self.assertIsNone(sheet.row_dimensions[9].height)
            self.assertFalse(sheet.row_dimensions[9].customHeight)
            self.assertTrue(all(sheet.cell(9, c).alignment.wrap_text for c in range(1, 7)))
            self.assertEqual(sheet.row_dimensions[8].height, 30)
            self.assertEqual(sheet.print_title_rows, "$1:$8")
            self.assertIn("$A$1:$F$9", sheet.print_area)
            self.assertEqual(sheet.page_setup.fitToWidth, 1)
            self.assertEqual(sheet.page_setup.fitToHeight, 0)

    def test_numbered_worksheets_only_replace_counters(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            library = root / "library"
            sheets = library / "drawing-sheets"
            sheets.mkdir(parents=True)
            source = sheets / "alex-generic-pcba.kicad_wks"
            original = b'\xef\xbb\xbf(kicad_wks (tbtext "${#}/${##} ${VARIANT} ${ProjectTitle}"))\r\n'
            source.write_bytes(original)
            work = root / "output"
            work.mkdir()
            with (patch.dict(os.environ, KICAD_LIB_ROOT=str(library),
                             JOBSET_OUTPUT_WORK_PATH=str(work)),
                  patch.object(formatter, "assembly_page_count", return_value=2)):
                formatter.prepare_assembly_worksheets()
                for page, view in ((1, "top"), (2, "bottom")):
                    self.assertEqual((work / f"_work/assembly-{view}.kicad_wks").read_bytes(),
                                     original.replace(b"${##}", b"2").replace(b"${#}", str(page).encode()))
                self.assertEqual(source.read_bytes(), original)
                # A stale or redirected target cannot be silently overwritten.
                with self.assertRaises(FileExistsError):
                    formatter.prepare_assembly_worksheets()

    def test_top_only_worksheet_uses_one_of_one(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            sheets = root / "library/drawing-sheets"
            sheets.mkdir(parents=True)
            source = sheets / "alex-generic-pcba.kicad_wks"
            source.write_bytes(b'(tbtext "${#}/${##}")')
            output = root / "output"
            output.mkdir()
            with (patch.dict(os.environ, KICAD_LIB_ROOT=str(root / "library"),
                             JOBSET_OUTPUT_WORK_PATH=str(output)),
                  patch.object(formatter, "assembly_page_count", return_value=1)):
                formatter.prepare_assembly_worksheets()
            self.assertEqual((output / "_work/assembly-top.kicad_wks").read_bytes(),
                             b'(tbtext "1/1")')
            self.assertFalse((output / "_work/assembly-bottom.kicad_wks").exists())

    def test_missing_canonical_worksheet_fails(self):
        with tempfile.TemporaryDirectory() as directory:
            with patch.dict(os.environ, KICAD_LIB_ROOT=directory, JOBSET_OUTPUT_WORK_PATH=directory):
                with self.assertRaisesRegex(FileNotFoundError, "KICAD_LIB_ROOT"):
                    formatter.prepare_assembly_worksheets()

    def test_rejects_unexpected_bom_schema(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "bom.csv"
            source.write_text("wrong,columns\n1,2\n")
            with self.assertRaisesRegex(ValueError, "Unexpected native BOM columns"):
                formatter.create_bom(source, Path(directory) / "out.xlsx", {}, "", datetime.now())

    def test_merge_preserves_page_order_and_dimensions(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for name, width in (("top", 100), ("bottom", 200)):
                writer = PdfWriter()
                writer.add_blank_page(width=width, height=300)
                writer.write(root / f"{name}.pdf")
            target = root / "assembly.pdf"
            formatter.merge_assembly(root / "top.pdf", root / "bottom.pdf", target, "Assembly")
            reader = PdfReader(target)
            self.assertEqual([page.mediabox.width for page in reader.pages], [100, 200])
            self.assertEqual([item.title for item in reader.outline], ["Top", "Bottom"])

    def test_top_only_assembly_is_published_without_rewriting(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=300)
            writer.add_metadata({"/Title": "Native top drawing"})
            top = root / "top.pdf"
            writer.write(top)
            target = root / "assembly.pdf"
            formatter.merge_assembly(top, None, target, "Ignored merge title")
            self.assertEqual(target.read_bytes(), top.read_bytes())
            reader = PdfReader(target)
            self.assertEqual(len(reader.pages), 1)
            self.assertEqual(reader.metadata.title, "Native top drawing")


if __name__ == "__main__":
    unittest.main()
